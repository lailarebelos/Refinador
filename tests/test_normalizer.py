import json
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from normalizer import (
    build_config,
    normalize,
    is_evaluable_for_sentiment,
    check_sentiment_eligibility,
    _infer_type_from_content,
    _auto_col_width,
    OutputColumn,
    _sanitize_short_name,
    _infer_value_map,
    _apply_value_map,
    _UF_MAP,
)


# ── Fixtures ──────────────────────────────────────────────────────────────────

def _write_sample_qsf(path: Path) -> None:
    payload = {
        "SurveyElements": [
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q1", "QuestionText": "Qual seu genero?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "Masculino"}, "2": {"Display": "Feminino"}},
                "ChoiceOrder": [1, 2],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q2", "QuestionText": "Conte um comentario sobre o atendimento.",
                "QuestionType": "TE", "Selector": "SL", "Choices": {}, "ChoiceOrder": [],
            }},
        ]
    }
    path.write_text(json.dumps(payload), encoding="utf-8")


def _write_sample_xlsx(path: Path) -> None:
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1"); ws.cell(1, 2, "Q2")
    ws.cell(2, 1, "Qual seu genero?"); ws.cell(2, 2, "Conte um comentario sobre o atendimento.")
    ws.cell(3, 1, "1"); ws.cell(3, 2, "Muito bom")
    ws.cell(4, 1, "2"); ws.cell(4, 2, "Precisa melhorar")
    wb.save(path)


def _write_qsf_with_timing(path: Path) -> None:
    payload = {
        "SurveyElements": [
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q1", "QuestionText": "Qual seu genero?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "Masculino"}, "2": {"Display": "Feminino"}},
                "ChoiceOrder": [1, 2],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q_TIMER", "QuestionText": "Page Timer",
                "QuestionType": "Timing", "Selector": "PageTimer",
                "Choices": {}, "ChoiceOrder": [],
            }},
        ]
    }
    path.write_text(json.dumps(payload), encoding="utf-8")


def _write_xlsx_with_timing(path: Path) -> None:
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1"); ws.cell(1, 2, "Q_TIMER_Page Submit")
    ws.cell(2, 1, "Qual seu genero?"); ws.cell(2, 2, "Page Submit")
    ws.cell(3, 1, "1"); ws.cell(3, 2, "5432")
    ws.cell(4, 1, "2"); ws.cell(4, 2, "3210")
    wb.save(path)


# ── Helpers de fabricação de respostas ────────────────────────────────────────

def _opinionated_series() -> pd.Series:
    return pd.Series([
        "O atendimento foi excelente, muito rápido e eficiente",
        "Precisa melhorar bastante o tempo de espera na fila",
        "Ótima experiência no geral, equipe muito atenciosa",
        "Funcionários educados mas o ambiente estava bagunçado",
        "Superou minhas expectativas, voltarei com certeza",
    ])


def _factual_city_series() -> pd.Series:
    return pd.Series(["São Paulo", "Rio de Janeiro", "Belo Horizonte",
                      "Curitiba", "Porto Alegre", "Salvador", "Fortaleza"])


def _factual_name_series() -> pd.Series:
    return pd.Series(["João Silva", "Maria Souza", "Pedro Oliveira",
                      "Ana Lima", "Carlos Mendes"])


# ══════════════════════════════════════════════════════════════════════════════
# 1. Testes de build_config / normalize (regressão)
# ══════════════════════════════════════════════════════════════════════════════

class BuildConfigTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        base = Path(self.tmpdir.name)
        self.qsf  = base / "s.qsf"
        self.xlsx = base / "e.xlsx"
        self.out  = base / "n.xlsx"
        _write_sample_qsf(self.qsf)
        _write_sample_xlsx(self.xlsx)

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_qsf_types_mapped_correctly(self):
        config, df = build_config(str(self.xlsx), str(self.qsf))
        by_id = {r["original_id"]: r for r in config}
        self.assertEqual(by_id["Q1"]["type"], "single_choice")
        self.assertEqual(by_id["Q2"]["type"], "open_text")
        self.assertEqual(by_id["Q2"]["alternatives"], "Texto livre")

    def test_normalize_produces_valid_workbook(self):
        config, df = build_config(str(self.xlsx), str(self.qsf))
        result = normalize(str(self.xlsx), str(self.out), config, df)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["rows"], 2)
        self.assertEqual(result["columns"], 2)
        wb = load_workbook(self.out)
        self.assertEqual(wb.sheetnames, ["Data", "Codebook"])
        ws = wb["Data"]
        # Q1 "Qual seu genero?" → only "Genero" survives stopword filter; no fallback padding
        self.assertEqual(ws["A1"].value, "Q1_Genero")
        # Q2 "Conte um comentario sobre o atendimento." → "sobre" is a preposition (now filtered),
        # "atendimento" is the actual topic — semantically correct new behavior
        self.assertEqual(ws["B1"].value, "Q2_Conte_Comentario_Atendimento")
        self.assertEqual(ws["A2"].value, 1)
        self.assertEqual(ws["B3"].value, "Precisa melhorar")

    def test_codebook_column_names(self):
        config, df = build_config(str(self.xlsx), str(self.qsf))
        normalize(str(self.xlsx), str(self.out), config, df)
        wb = load_workbook(self.out)
        wc = wb["Codebook"]
        headers = [wc.cell(1, c).value for c in range(1, 7)]
        self.assertEqual(headers[0], "question_code")
        self.assertEqual(headers[1], "question")

    def test_nps_alternatives_portuguese(self):
        """NPS group column must show Portuguese labels."""
        tmpdir = tempfile.TemporaryDirectory()
        base = Path(tmpdir.name)
        qsf_path = base / "nps.qsf"
        xlsx_path = base / "nps.xlsx"
        out_path  = base / "nps_out.xlsx"

        payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
            "DataExportTag": "Q1", "QuestionText": "NPS",
            "QuestionType": "MC", "Selector": "NPS",
            "Choices": {str(i): {"Display": str(i)} for i in range(11)},
            "ChoiceOrder": list(range(11)),
        }}]}
        qsf_path.write_text(json.dumps(payload), encoding="utf-8")

        from openpyxl import Workbook as WB
        wb2 = WB(); ws2 = wb2.active
        ws2.cell(1, 1, "Q1_NPS_GROUP")
        ws2.cell(2, 1, "NPS group")
        for r, v in enumerate([1, 2, 3], start=3):
            ws2.cell(r, 1, str(v))
        wb2.save(xlsx_path)

        config, df = build_config(str(xlsx_path), str(qsf_path))
        nps_cfg = next(c for c in config if "NPS_GROUP" in c["original_id"])
        self.assertIn("Detrator", nps_cfg["alternatives"])
        self.assertIn("Promotor", nps_cfg["alternatives"])
        self.assertNotIn("Promoter", nps_cfg["alternatives"])
        tmpdir.cleanup()

    def test_short_name_no_padding_words(self):
        """Short names must not be padded with 'Perfil', 'Respondente' or 'Pesquisa'."""
        name = _sanitize_short_name("Qual seu gênero?")
        for bad in ("Perfil", "Respondente", "Pesquisa"):
            self.assertNotIn(bad, name)

    def test_timing_columns_excluded(self):
        tmpdir = tempfile.TemporaryDirectory()
        base = Path(tmpdir.name)
        qsf = base / "t.qsf"; xlsx = base / "t.xlsx"
        _write_qsf_with_timing(qsf)
        _write_xlsx_with_timing(xlsx)
        config, _ = build_config(str(xlsx), str(qsf))
        ids = [r["original_id"] for r in config]
        self.assertIn("Q1", ids)
        self.assertNotIn("Q_TIMER_Page Submit", ids)
        tmpdir.cleanup()


# ══════════════════════════════════════════════════════════════════════════════
# 2. is_evaluable_for_sentiment  (verifica QUANTIDADE de texto)
# ══════════════════════════════════════════════════════════════════════════════

class EvaluableTests(unittest.TestCase):
    def test_enough_real_text_is_evaluable(self):
        self.assertTrue(is_evaluable_for_sentiment(_opinionated_series()))

    def test_all_empty_not_evaluable(self):
        self.assertFalse(is_evaluable_for_sentiment(pd.Series([None, None, None])))

    def test_trivial_values_not_evaluable(self):
        self.assertFalse(is_evaluable_for_sentiment(pd.Series(["N/A", "sim", "não", "n/a", "sim"])))

    def test_too_few_rows_not_evaluable(self):
        self.assertFalse(is_evaluable_for_sentiment(pd.Series(["Ótimo atendimento", "Muito ruim mesmo"])))

    def test_very_short_texts_not_evaluable(self):
        self.assertFalse(is_evaluable_for_sentiment(pd.Series(["ok", "ok", "ok", "ok"])))


# ══════════════════════════════════════════════════════════════════════════════
# 3. check_sentiment_eligibility  — exemplos EXATOS do usuário
#    Testa a SEPARAÇÃO entre tipo estrutural e elegibilidade de sentimento.
# ══════════════════════════════════════════════════════════════════════════════

class SentimentEligibilityTests(unittest.TestCase):

    # ── Perguntas ELEGÍVEIS ────────────────────────────────────────────────────

    def test_comente_experiencia_eligible(self):
        e = check_sentiment_eligibility("Comente sua experiência")
        self.assertTrue(e.eligible, e.summary())

    def test_o_que_podemos_melhorar_eligible(self):
        e = check_sentiment_eligibility("O que poderíamos melhorar?")
        self.assertTrue(e.eligible, e.summary())

    def test_por_que_deu_essa_nota_eligible(self):
        e = check_sentiment_eligibility("Por que deu essa nota?")
        self.assertTrue(e.eligible, e.summary())

    def test_satisfacao_geral_eligible(self):
        e = check_sentiment_eligibility("Como você avalia sua satisfação com o atendimento?")
        self.assertTrue(e.eligible, e.summary())

    def test_sugestao_eligible(self):
        e = check_sentiment_eligibility("Tem alguma sugestão para melhorarmos?")
        self.assertTrue(e.eligible, e.summary())

    def test_feedback_geral_eligible(self):
        e = check_sentiment_eligibility("Deixe seu feedback sobre o serviço")
        self.assertTrue(e.eligible, e.summary())

    def test_reclamacao_eligible(self):
        e = check_sentiment_eligibility("Registre aqui sua reclamação ou elogio")
        self.assertTrue(e.eligible, e.summary())

    def test_opiniao_eligible(self):
        e = check_sentiment_eligibility("Qual a sua opinião sobre o produto?")
        self.assertTrue(e.eligible, e.summary())

    # ── Perguntas NÃO ELEGÍVEIS ───────────────────────────────────────────────

    def test_outros_qual_empresa_not_eligible(self):
        e = check_sentiment_eligibility("Outros: qual empresa?")
        self.assertFalse(e.eligible, e.summary())

    def test_nome_da_empresa_not_eligible(self):
        e = check_sentiment_eligibility("Digite o nome da empresa")
        self.assertFalse(e.eligible, e.summary())

    def test_qual_cidade_not_eligible(self):
        e = check_sentiment_eligibility("Qual cidade?")
        self.assertFalse(e.eligible, e.summary())

    def test_nome_do_cliente_not_eligible(self):
        e = check_sentiment_eligibility("Nome do cliente")
        self.assertFalse(e.eligible, e.summary())

    def test_email_not_eligible(self):
        e = check_sentiment_eligibility("Informe seu e-mail")
        self.assertFalse(e.eligible, e.summary())

    def test_telefone_not_eligible(self):
        e = check_sentiment_eligibility("Telefone para contato")
        self.assertFalse(e.eligible, e.summary())

    def test_cargo_not_eligible(self):
        e = check_sentiment_eligibility("Qual o seu cargo?")
        # "cargo" é cadastral mesmo que "qual" possa sugerir pergunta
        self.assertFalse(e.eligible, e.summary())

    def test_especifique_not_eligible(self):
        e = check_sentiment_eligibility("Especifique qual produto")
        self.assertFalse(e.eligible, e.summary())

    # ── Elegibilidade com série: conteúdo reforça o resultado ─────────────────

    def test_opinion_question_with_opinionated_content_eligible(self):
        e = check_sentiment_eligibility("Comente sua experiência", _opinionated_series())
        self.assertTrue(e.eligible, e.summary())

    def test_factual_question_with_city_content_not_eligible(self):
        e = check_sentiment_eligibility("Qual cidade você mora?", _factual_city_series())
        self.assertFalse(e.eligible, e.summary())

    def test_factual_question_with_names_not_eligible(self):
        e = check_sentiment_eligibility("Nome do respondente", _factual_name_series())
        self.assertFalse(e.eligible, e.summary())

    def test_opinion_question_with_empty_series_not_eligible(self):
        e = check_sentiment_eligibility("Comente sua experiência", pd.Series([None, None, None]))
        self.assertFalse(e.eligible, e.summary())

    # ── Score e reasons são informativos ──────────────────────────────────────

    def test_eligible_question_has_positive_score(self):
        e = check_sentiment_eligibility("Comente sua experiência")
        self.assertGreater(e.score, 0)

    def test_factual_question_has_negative_or_zero_score(self):
        e = check_sentiment_eligibility("Nome do cliente")
        self.assertLess(e.score, 1.5)

    def test_summary_contains_label(self):
        e = check_sentiment_eligibility("Comente sua experiência")
        self.assertIn("Elegível", e.summary())

    def test_summary_contains_reasons(self):
        e = check_sentiment_eligibility("Comente sua experiência")
        self.assertGreater(len(e.reasons_for), 0)

    def test_not_eligible_has_reasons_against(self):
        e = check_sentiment_eligibility("Digite o nome da empresa")
        self.assertGreater(len(e.reasons_against), 0)


# ══════════════════════════════════════════════════════════════════════════════
# 4. _infer_type_from_content  (tipo ESTRUTURAL — separado de sentimento)
# ══════════════════════════════════════════════════════════════════════════════

class InferTypeTests(unittest.TestCase):

    def test_long_text_is_open_text(self):
        s = pd.Series([
            "O atendimento foi excelente e o funcionário foi muito prestativo na resolução",
            "Achei o ambiente muito agradável e bem organizado para atender os clientes",
            "Precisaria melhorar o tempo de espera na fila que foi bem longo para mim",
        ])
        self.assertEqual(_infer_type_from_content(s), "open_text")

    def test_city_names_are_open_text_structurally(self):
        # Factual, mas estruturalmente é open_text — sentimento bloqueado por check_sentiment_eligibility
        s = pd.Series(["São Paulo", "Rio de Janeiro", "Belo Horizonte",
                       "Curitiba", "Porto Alegre", "Salvador", "Fortaleza"])
        self.assertEqual(_infer_type_from_content(s), "open_text")

    def test_nps_range_detected(self):
        s = pd.Series(["0", "3", "5", "7", "8", "9", "10", "6", "4", "10"])
        self.assertEqual(_infer_type_from_content(s), "nps")

    def test_few_unique_numeric_is_single_choice(self):
        s = pd.Series(["1", "2", "3", "1", "2", "3", "1"])
        self.assertEqual(_infer_type_from_content(s), "single_choice")

    def test_few_unique_text_is_single_choice(self):
        s = pd.Series(["Masculino", "Feminino", "Masculino", "Feminino", "Masculino"])
        self.assertEqual(_infer_type_from_content(s), "single_choice")

    def test_too_few_values_returns_none(self):
        self.assertIsNone(_infer_type_from_content(pd.Series(["abc"])))

    def test_high_numeric_range_is_numeric(self):
        s = pd.Series([str(i * 100) for i in range(1, 20)])
        self.assertEqual(_infer_type_from_content(s), "numeric")

    def test_long_categorical_labels_are_single_choice(self):
        """Income brackets have long labels but are a finite closed set — must be single_choice."""
        labels = [
            "Até 2 salários-mínimos (até R$ 3.018,00)",
            "De 2 a 4 salários-mínimos (de R$ 3.018,01 a R$ 6.036,00)",
            "De 4 a 10 salários-mínimos (de R$ 6.036,01 a R$15.090,00)",
            "De 10 a 20 salários-mínimos (de R$15.090,01 a R$ 30.180,00)",
            "Mais de 20 salários-mínimos (R$ 30.180,01 ou mais)",
        ]
        # 5 unique labels repeated across 40 respondents → avg_reps = 8
        s = pd.Series(labels * 8)
        self.assertEqual(_infer_type_from_content(s), "single_choice")

    def test_long_categorical_labels_encoded_to_numbers(self):
        """After full normalize(), income bracket text must become int in Data tab."""
        import tempfile
        from openpyxl import Workbook as WB
        tmpdir = tempfile.TemporaryDirectory()
        base = Path(tmpdir.name)
        qsf_path = base / "q.qsf"
        xlsx_path = base / "d.xlsx"
        out_path  = base / "o.xlsx"

        labels = [
            "Até 2 salários-mínimos",
            "De 2 a 4 salários-mínimos",
            "De 4 a 10 salários-mínimos",
            "De 10 a 20 salários-mínimos",
            "Mais de 20 salários-mínimos",
        ]
        payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
            "DataExportTag": "Q1", "QuestionText": "Qual sua faixa de renda?",
            "QuestionType": "MC", "Selector": "SAVR",
            "Choices": {str(i + 1): {"Display": lbl} for i, lbl in enumerate(labels)},
            "ChoiceOrder": list(range(1, 6)),
        }}]}
        qsf_path.write_text(json.dumps(payload), encoding="utf-8")

        wb2 = WB(); ws2 = wb2.active
        ws2.cell(1, 1, "Q1")
        ws2.cell(2, 1, "Qual sua faixa de renda?")
        for r, lbl in enumerate(labels * 2, start=3):
            ws2.cell(r, 1, lbl)
        wb2.save(xlsx_path)

        config, df = build_config(str(xlsx_path), str(qsf_path))
        normalize(str(xlsx_path), str(out_path), config, df)
        wb_out = load_workbook(out_path)
        ws_out = wb_out["Data"]
        data_values = [ws_out.cell(r, 1).value for r in range(2, 12)]
        for v in data_values:
            self.assertIsInstance(v, int, f"Expected int but got {v!r}")
        tmpdir.cleanup()

    def test_operational_metadata_excluded(self):
        """ResponseId, IPAddress and similar columns must not appear in config."""
        tmpdir = tempfile.TemporaryDirectory()
        base = Path(tmpdir.name)
        qsf_path = base / "q.qsf"
        xlsx_path = base / "d.xlsx"

        payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
            "DataExportTag": "Q1", "QuestionText": "Gênero",
            "QuestionType": "MC", "Selector": "SAVR",
            "Choices": {"1": {"Display": "Masculino"}, "2": {"Display": "Feminino"}},
            "ChoiceOrder": [1, 2],
        }}]}
        qsf_path.write_text(json.dumps(payload), encoding="utf-8")

        from openpyxl import Workbook as WB
        wb2 = WB(); ws2 = wb2.active
        ws2.cell(1, 1, "Q1"); ws2.cell(1, 2, "ResponseId"); ws2.cell(1, 3, "IPAddress")
        ws2.cell(2, 1, "Gênero"); ws2.cell(2, 2, "Response ID"); ws2.cell(2, 3, "IP Address")
        for r in range(3, 13):
            ws2.cell(r, 1, "1" if r % 2 else "2")
            ws2.cell(r, 2, f"R_{r:04d}ABCDEF1234567890")
            ws2.cell(r, 3, f"192.168.{r}.1")
        wb2.save(xlsx_path)

        config, _ = build_config(str(xlsx_path), str(qsf_path))
        ids = [c["original_id"] for c in config]
        self.assertIn("Q1", ids)
        self.assertNotIn("ResponseId", ids)
        self.assertNotIn("IPAddress", ids)
        tmpdir.cleanup()


# ══════════════════════════════════════════════════════════════════════════════
# 5. _auto_col_width
# ══════════════════════════════════════════════════════════════════════════════

class AutoColWidthTests(unittest.TestCase):

    def test_open_text_width_capped_at_55(self):
        col = OutputColumn("H", "Q", "open_text", ["x" * 300] * 10)
        self.assertLessEqual(_auto_col_width(col), 55)
        self.assertGreaterEqual(_auto_col_width(col), 10)

    def test_other_width_capped_at_30(self):
        col = OutputColumn("H_" + "x" * 100, "Q", "single_choice", ["1"])
        self.assertLessEqual(_auto_col_width(col), 30)

    def test_short_content_still_meets_minimum(self):
        col = OutputColumn("Q1", "Q", "single_choice", ["1", "2", "1"])
        self.assertGreaterEqual(_auto_col_width(col), 10)


# ══════════════════════════════════════════════════════════════════════════════
# 6. EncodingTests — new encoding strategies
# ══════════════════════════════════════════════════════════════════════════════

def _write_text_single_choice_xlsx(path: Path) -> None:
    """XLSX with text values 'Sim'/'Não' for Q1."""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1")
    ws.cell(2, 1, "Você recomendaria?")
    ws.cell(3, 1, "Sim")
    ws.cell(4, 1, "Não")
    ws.cell(5, 1, "Sim")
    wb.save(path)


def _write_text_single_choice_qsf(path: Path) -> None:
    payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
        "DataExportTag": "Q1", "QuestionText": "Você recomendaria?",
        "QuestionType": "MC", "Selector": "SAVR",
        "Choices": {"1": {"Display": "Sim"}, "2": {"Display": "Não"}},
        "ChoiceOrder": [1, 2],
    }}]}
    path.write_text(json.dumps(payload), encoding="utf-8")


def _write_satisfaction_xlsx(path: Path) -> None:
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1")
    ws.cell(2, 1, "Como avalia o atendimento?")
    for r, v in enumerate(["Muito insatisfeito", "Insatisfeito", "Neutro",
                            "Satisfeito", "Muito satisfeito"], start=3):
        ws.cell(r, 1, v)
    wb.save(path)


def _write_satisfaction_qsf(path: Path) -> None:
    payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
        "DataExportTag": "Q1", "QuestionText": "Como avalia o atendimento?",
        "QuestionType": "MC", "Selector": "SAVR",
        "Choices": {
            "1": {"Display": "Muito insatisfeito"},
            "2": {"Display": "Insatisfeito"},
            "3": {"Display": "Neutro"},
            "4": {"Display": "Satisfeito"},
            "5": {"Display": "Muito satisfeito"},
        },
        "ChoiceOrder": [1, 2, 3, 4, 5],
    }}]}
    path.write_text(json.dumps(payload), encoding="utf-8")


def _write_onehot_xlsx(path: Path) -> None:
    """Four one-hot sub-columns Q1_1…Q1_4 where exactly one is selected per row."""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1_1"); ws.cell(1, 2, "Q1_2")
    ws.cell(1, 3, "Q1_3"); ws.cell(1, 4, "Q1_4")
    ws.cell(2, 1, "Cor preferida - Vermelho"); ws.cell(2, 2, "Cor preferida - Verde")
    ws.cell(2, 3, "Cor preferida - Azul");     ws.cell(2, 4, "Cor preferida - Amarelo")
    # Row 1: chose option 1
    ws.cell(3, 1, "1"); ws.cell(3, 2, ""); ws.cell(3, 3, ""); ws.cell(3, 4, "")
    # Row 2: chose option 3
    ws.cell(4, 1, ""); ws.cell(4, 2, ""); ws.cell(4, 3, "1"); ws.cell(4, 4, "")
    # Row 3: chose option 2
    ws.cell(5, 1, ""); ws.cell(5, 2, "1"); ws.cell(5, 3, ""); ws.cell(5, 4, "")
    wb.save(path)


def _write_onehot_qsf(path: Path) -> None:
    payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
        "DataExportTag": "Q1", "QuestionText": "Qual sua cor preferida?",
        "QuestionType": "MC", "Selector": "SAVR",
        "Choices": {
            "1": {"Display": "Vermelho"},
            "2": {"Display": "Verde"},
            "3": {"Display": "Azul"},
            "4": {"Display": "Amarelo"},
        },
        "ChoiceOrder": [1, 2, 3, 4],
    }}]}
    path.write_text(json.dumps(payload), encoding="utf-8")


def _write_mavr_xlsx(path: Path) -> None:
    """Three MAVR binary sub-columns."""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1_1"); ws.cell(1, 2, "Q1_2"); ws.cell(1, 3, "Q1_3")
    ws.cell(2, 1, "Produto A"); ws.cell(2, 2, "Produto B"); ws.cell(2, 3, "Produto C")
    ws.cell(3, 1, "1"); ws.cell(3, 2, "0"); ws.cell(3, 3, "1")
    ws.cell(4, 1, "0"); ws.cell(4, 2, "1"); ws.cell(4, 3, "0")
    wb.save(path)


def _write_mavr_qsf(path: Path) -> None:
    payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
        "DataExportTag": "Q1", "QuestionText": "Quais produtos você usa?",
        "QuestionType": "MC", "Selector": "MAVR",
        "Choices": {
            "1": {"Display": "Produto A"},
            "2": {"Display": "Produto B"},
            "3": {"Display": "Produto C"},
        },
        "ChoiceOrder": [1, 2, 3],
    }}]}
    path.write_text(json.dumps(payload), encoding="utf-8")


def _write_open_text_xlsx(path: Path) -> None:
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1")
    ws.cell(2, 1, "Deixe seu comentário.")
    ws.cell(3, 1, "Excelente serviço")
    ws.cell(4, 1, "Poderia melhorar")
    ws.cell(5, 1, "")
    wb.save(path)


def _write_open_text_qsf(path: Path) -> None:
    payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
        "DataExportTag": "Q1", "QuestionText": "Deixe seu comentário.",
        "QuestionType": "TE", "Selector": "SL", "Choices": {}, "ChoiceOrder": [],
    }}]}
    path.write_text(json.dumps(payload), encoding="utf-8")


def _write_uf_xlsx(path: Path) -> None:
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "Q1")
    ws.cell(2, 1, "Qual seu estado?")
    ws.cell(3, 1, "São Paulo")
    ws.cell(4, 1, "Minas Gerais")
    ws.cell(5, 1, "Rio de Janeiro")
    wb.save(path)


def _write_uf_qsf(path: Path) -> None:
    payload = {"SurveyElements": [{"Element": "SQ", "Payload": {
        "DataExportTag": "Q1", "QuestionText": "Qual seu estado?",
        "QuestionType": "TE", "Selector": "SL", "Choices": {}, "ChoiceOrder": [],
    }}]}
    path.write_text(json.dumps(payload), encoding="utf-8")


class EncodingTests(unittest.TestCase):

    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.base = Path(self.tmpdir.name)

    def tearDown(self):
        self.tmpdir.cleanup()

    # ── helpers ───────────────────────────────────────────────────────────────

    def _run(self, write_qsf, write_xlsx):
        qsf = self.base / "s.qsf"
        xlsx = self.base / "e.xlsx"
        out = self.base / "n.xlsx"
        write_qsf(qsf)
        write_xlsx(xlsx)
        config, df = build_config(str(xlsx), str(qsf))
        result = normalize(str(xlsx), str(out), config, df)
        wb = load_workbook(out)
        return config, df, result, wb

    # ── tests ─────────────────────────────────────────────────────────────────

    def test_text_single_choice_encoded_to_numeric(self):
        """When XLSX has text values 'Sim'/'Não', Data must contain integers."""
        config, df, result, wb = self._run(
            _write_text_single_choice_qsf, _write_text_single_choice_xlsx)
        ws = wb["Data"]
        # All data cells in column A should be numeric (int), not text
        data_values = [ws.cell(r, 1).value for r in range(2, 5)]
        for v in data_values:
            self.assertIsNotNone(v)
            self.assertIsInstance(v, int, f"Expected int, got {type(v).__name__}: {v!r}")

    def test_satisfaction_scale_ordered_correctly(self):
        """Muito Insatisfeito=1, Insatisfeito=2, ..., Muito Satisfeito=5."""
        series = pd.Series(["Muito insatisfeito", "Insatisfeito", "Neutro",
                            "Satisfeito", "Muito satisfeito"])
        value_map, _ = _infer_value_map(series, None)
        from normalizer import _norm
        self.assertEqual(value_map[_norm("Muito insatisfeito")], 1)
        self.assertEqual(value_map[_norm("Muito satisfeito")], 5)
        self.assertLess(value_map[_norm("Insatisfeito")], value_map[_norm("Satisfeito")])

    def test_onehot_collapses_to_single_column(self):
        """Four SAVR one-hot sub-columns collapse to one numeric column."""
        config, df, result, wb = self._run(_write_onehot_qsf, _write_onehot_xlsx)
        ws = wb["Data"]
        # Should produce exactly 1 output column (the collapse host), not 4
        self.assertEqual(result["columns"], 1)
        # Values should be the choice codes (1, 3, 2) matching row selections
        data_values = [ws.cell(r, 1).value for r in range(2, 5)]
        self.assertEqual(data_values, [1, 3, 2])

    def test_multiple_choice_stays_as_indicators(self):
        """MAVR sub-columns become 0/1 indicator columns (3 in, 3 out)."""
        config, df, result, wb = self._run(_write_mavr_qsf, _write_mavr_xlsx)
        ws = wb["Data"]
        self.assertEqual(result["columns"], 3)
        # All data values must be 0 or 1
        for ci in range(1, 4):
            for ri in range(2, 4):
                v = ws.cell(ri, ci).value
                self.assertIn(v, (0, 1), f"cell({ri},{ci})={v!r}")

    def test_alternatives_filled_for_single_choice(self):
        """Codebook alternatives is not empty for single_choice column."""
        config, df, result, wb = self._run(
            _write_text_single_choice_qsf, _write_text_single_choice_xlsx)
        cfg = next(c for c in config if c["original_id"] == "Q1")
        self.assertTrue(cfg["alternatives"], "alternatives should not be empty")

    def test_uf_encoded_to_number(self):
        """São Paulo → 25, Minas Gerais → 13 via _UF_MAP."""
        from normalizer import _norm
        self.assertEqual(_UF_MAP[_norm("São Paulo")], 25)
        self.assertEqual(_UF_MAP[_norm("Minas Gerais")], 13)

    def test_open_text_preserved(self):
        """Open text values are NOT converted to numbers."""
        config, df, result, wb = self._run(_write_open_text_qsf, _write_open_text_xlsx)
        ws = wb["Data"]
        v = ws.cell(2, 1).value
        self.assertIsInstance(v, str, f"Expected str, got {type(v).__name__}: {v!r}")

    def test_data_codebook_consistency(self):
        """Every header in Data has exactly one row in Codebook."""
        config, df, result, wb = self._run(
            _write_text_single_choice_qsf, _write_text_single_choice_xlsx)
        ws_data = wb["Data"]
        ws_cb = wb["Codebook"]
        data_headers = [ws_data.cell(1, c).value for c in range(1, result["columns"] + 1)]
        cb_codes = [ws_cb.cell(r, 1).value for r in range(2, result["columns"] + 2)]
        self.assertEqual(sorted(data_headers), sorted(cb_codes))

    def test_infer_value_map_uses_qsf_choices(self):
        """_infer_value_map returns QSF-derived map when choices are available."""
        from normalizer import QsfQuestion
        qsf_q = QsfQuestion(
            export_tag="Q1", question_text="Gênero",
            question_type="MC", selector="SAVR",
            choices={1: "Masculino", 2: "Feminino"},
            has_text_entry=False, choice_order=[1, 2],
        )
        series = pd.Series(["Masculino", "Feminino", "Masculino"])
        value_map, alts = _infer_value_map(series, qsf_q)
        from normalizer import _norm
        self.assertEqual(value_map[_norm("Masculino")], 1)
        self.assertEqual(value_map[_norm("Feminino")], 2)
        self.assertIn("Masculino", alts)

    def test_apply_value_map_handles_missing_with_numeric_fallback(self):
        """_apply_value_map keeps numeric strings that aren't in the map."""
        result = _apply_value_map("3", {"sim": 1, "nao": 2})
        self.assertEqual(result, 3)

    def test_apply_value_map_returns_none_for_empty(self):
        self.assertIsNone(_apply_value_map("", {"a": 1}))
        self.assertIsNone(_apply_value_map(None, {"a": 1}))

    def test_subsumed_columns_excluded_by_default(self):
        """One-hot sub-columns that are subsumed get include=False."""
        qsf = self.base / "s.qsf"
        xlsx = self.base / "e.xlsx"
        _write_onehot_qsf(qsf)
        _write_onehot_xlsx(xlsx)
        config, _ = build_config(str(xlsx), str(qsf))
        subsumed = [c for c in config if c.get("_encoding_strategy") == "subsumed"]
        for c in subsumed:
            self.assertFalse(c["include"],
                             f"{c['original_id']} subsumed but include=True")

    def test_encoding_strategy_present_in_config(self):
        """All config rows must have _encoding_strategy set."""
        qsf = self.base / "s.qsf"
        xlsx = self.base / "e.xlsx"
        _write_text_single_choice_qsf(qsf)
        _write_text_single_choice_xlsx(xlsx)
        config, _ = build_config(str(xlsx), str(qsf))
        for row in config:
            self.assertIn("_encoding_strategy", row,
                          f"Missing _encoding_strategy in {row['original_id']}")
            self.assertIn(row["_encoding_strategy"],
                          ("direct", "text_encode", "onehot_collapse",
                           "indicator", "text_preserve", "subsumed"))


# ── Fixtures para testes de alinhamento ──────────────────────────────────────

def _write_alignment_qsf(path: Path) -> None:
    """QSF com 8 perguntas demográficas.
    Q4, Q7 e Q8 têm cabeçalhos de seção embutidos no QuestionText
    (padrão comum no Qualtrics). O XLSX exporta apenas o texto limpo.
    O pipeline deve gerar question_codes a partir do texto limpo do XLSX,
    NÃO dos cabeçalhos do QSF.
    """
    payload = {
        "SurveyElements": [
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q1",
                "QuestionText": "Em qual regiao do Brasil voce mora?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "Norte"}, "2": {"Display": "Sudeste"}, "3": {"Display": "Sul"}},
                "ChoiceOrder": [1, 2, 3],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q2",
                "QuestionText": "Em qual estado voce mora?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "SP"}, "2": {"Display": "RJ"}, "3": {"Display": "MG"}},
                "ChoiceOrder": [1, 2, 3],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q3",
                "QuestionText": "Em qual cidade voce mora?",
                "QuestionType": "TE", "Selector": "SL", "Choices": {}, "ChoiceOrder": [],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q4",
                # Cabeçalho de seção embutido: "Perfil do motorista."
                "QuestionText": "Perfil do motorista. Com qual genero voce se identifica?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "Feminino"}, "2": {"Display": "Masculino"}, "3": {"Display": "Outro"}},
                "ChoiceOrder": [1, 2, 3],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q5",
                "QuestionText": "Qual e a sua faixa etaria?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "18-24"}, "2": {"Display": "25-34"}, "3": {"Display": "35+"}},
                "ChoiceOrder": [1, 2, 3],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q6",
                "QuestionText": "Qual e o seu nivel de formacao academica?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "Medio"}, "2": {"Display": "Superior"}, "3": {"Display": "Pos"}},
                "ChoiceOrder": [1, 2, 3],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q7",
                # Cabeçalho de seção embutido: "Formacao e renda."
                "QuestionText": "Formacao e renda. Somando a sua renda com a renda das pessoas que moram com voce, qual e a renda familiar mensal?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "Ate 1500"}, "2": {"Display": "1500-3000"}, "3": {"Display": "Acima 3000"}},
                "ChoiceOrder": [1, 2, 3],
            }},
            {"Element": "SQ", "Payload": {
                "DataExportTag": "Q8",
                # Cabeçalho de seção embutido: "Estado civil."
                "QuestionText": "Estado civil. Voce possui filhos?",
                "QuestionType": "MC", "Selector": "SAVR",
                "Choices": {"1": {"Display": "Sim"}, "2": {"Display": "Nao"}},
                "ChoiceOrder": [1, 2],
            }},
        ]
    }
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def _write_alignment_xlsx(path: Path) -> None:
    """XLSX com texto limpo em row 2 — sem cabeçalhos de seção."""
    wb = Workbook(); ws = wb.active
    questions = [
        ("Q1", "Em qual regiao do Brasil voce mora?", "1", "3"),
        ("Q2", "Em qual estado voce mora?", "2", "3"),
        ("Q3", "Em qual cidade voce mora?", "Sao Paulo", "Curitiba"),
        ("Q4", "Com qual genero voce se identifica?", "1", "2"),
        ("Q5", "Qual e a sua faixa etaria?", "1", "2"),
        ("Q6", "Qual e o seu nivel de formacao academica?", "2", "3"),
        ("Q7", "Somando a sua renda com a renda das pessoas que moram com voce, qual e a renda familiar mensal?", "1", "2"),
        ("Q8", "Voce possui filhos?", "1", "2"),
    ]
    for col_idx, (qid, qtxt, v1, v2) in enumerate(questions, 1):
        ws.cell(1, col_idx, qid)
        ws.cell(2, col_idx, qtxt)
        ws.cell(3, col_idx, v1)
        ws.cell(4, col_idx, v2)
    wb.save(path)


# ══════════════════════════════════════════════════════════════════════════════
# Testes de alinhamento question_code ↔ question_text
# ══════════════════════════════════════════════════════════════════════════════

class AlignmentTests(unittest.TestCase):
    """Valida que question_code nunca é associado à pergunta errada.

    Testa o bug de 'capítulo embutido': quando o QSF QuestionText de uma
    pergunta começa com um cabeçalho de seção (ex.: 'Perfil do motorista.
    Com qual gênero...'), o question_code gerado deve refletir o texto limpo
    da planilha XLSX (row 2), NÃO o cabeçalho do QSF.
    """

    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        base = Path(self.tmpdir.name)
        self.qsf  = base / "align.qsf"
        self.xlsx = base / "align.xlsx"
        self.out  = base / "align_out.xlsx"
        _write_alignment_qsf(self.qsf)
        _write_alignment_xlsx(self.xlsx)

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_short_name_not_polluted_by_qsf_chapter_headings(self):
        """short_name de cada pergunta NÃO deve conter palavras do cabeçalho
        de seção embutido no QSF — apenas do texto real da pergunta (XLSX)."""
        config, _ = build_config(str(self.xlsx), str(self.qsf))
        by_id = {r["original_id"]: r for r in config}

        q4 = by_id["Q4"]
        self.assertNotIn("perfil", q4["short_name"].lower(),
            f"Q4 (gênero): short_name '{q4['short_name']}' contém cabeçalho QSF 'perfil'")
        self.assertNotIn("motorista", q4["short_name"].lower(),
            f"Q4 (gênero): short_name '{q4['short_name']}' contém cabeçalho QSF 'motorista'")

        q7 = by_id["Q7"]
        self.assertNotIn("formacao", q7["short_name"].lower(),
            f"Q7 (renda): short_name '{q7['short_name']}' contém cabeçalho QSF 'formacao'")

        q8 = by_id["Q8"]
        self.assertNotIn("civil", q8["short_name"].lower(),
            f"Q8 (filhos): short_name '{q8['short_name']}' contém cabeçalho QSF 'civil'")

    def test_config_original_id_matches_source(self):
        """original_id deve corresponder ao qualtrics_id da coluna de origem."""
        config, df = build_config(str(self.xlsx), str(self.qsf))
        for row in config:
            qid = row["original_id"]
            self.assertIn(qid, df.columns,
                f"original_id '{qid}' não encontrado no DataFrame")

    def test_codebook_code_paired_with_correct_question(self):
        """No Codebook exportado: question_code em coluna A deve estar emparelhado
        com a question_text correta em coluna B — nunca por índice, sempre por ID."""
        config, df = build_config(str(self.xlsx), str(self.qsf))
        normalize(str(self.xlsx), str(self.out), config, df)
        wb = load_workbook(self.out)
        ws = wb["Codebook"]

        code_to_text: dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                code_to_text[str(row[0])] = str(row[1])

        def find_by_prefix(prefix: str) -> tuple[str, str]:
            for code, text in code_to_text.items():
                if code.startswith(prefix + "_") or code == prefix:
                    return code, text
            return "", ""

        q4_code, q4_text = find_by_prefix("Q4")
        self.assertTrue(q4_code, "Q4 não encontrado no Codebook")
        self.assertIn("genero", q4_text.lower(),
            f"Q4 no Codebook: code='{q4_code}' emparelhado com pergunta errada: '{q4_text}'")

        q7_code, q7_text = find_by_prefix("Q7")
        self.assertTrue(q7_code, "Q7 não encontrado no Codebook")
        self.assertIn("renda", q7_text.lower(),
            f"Q7 no Codebook: code='{q7_code}' emparelhado com pergunta errada: '{q7_text}'")

        q8_code, q8_text = find_by_prefix("Q8")
        self.assertTrue(q8_code, "Q8 não encontrado no Codebook")
        self.assertIn("filhos", q8_text.lower(),
            f"Q8 no Codebook: code='{q8_code}' emparelhado com pergunta errada: '{q8_text}'")

    def test_data_tab_codes_align_with_codebook(self):
        """Todo header da aba Data deve ter linha correspondente no Codebook
        com o mesmo question_code. Valida que não há pareamento por índice."""
        config, df = build_config(str(self.xlsx), str(self.qsf))
        normalize(str(self.xlsx), str(self.out), config, df)
        wb = load_workbook(self.out)
        ws_data = wb["Data"]
        ws_cb   = wb["Codebook"]

        data_headers = {ws_data.cell(1, c).value
                        for c in range(1, ws_data.max_column + 1)
                        if ws_data.cell(1, c).value}
        cb_codes = {ws_cb.cell(r, 1).value
                    for r in range(2, ws_cb.max_row + 1)
                    if ws_cb.cell(r, 1).value}

        missing = data_headers - cb_codes
        self.assertFalse(missing,
            f"Codes presentes na aba Data mas ausentes no Codebook: {missing}")

    def test_no_duplicate_question_codes(self):
        """Não podem existir question_codes duplicados na aba Data."""
        config, df = build_config(str(self.xlsx), str(self.qsf))
        normalize(str(self.xlsx), str(self.out), config, df)
        wb = load_workbook(self.out)
        ws = wb["Data"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)
                   if ws.cell(1, c).value]
        self.assertEqual(len(headers), len(set(headers)),
            f"question_codes duplicados na aba Data: {headers}")


if __name__ == "__main__":
    unittest.main()
